# -*- coding: utf-8 -*-
"""
운정중앙역 주변 아파트 전·월세 매물 수집기 (네이버 부동산 기준)

- 운정중앙역 반경 RADIUS_KM 안의 아파트 단지를 찾고
- 전세(INCLUDE_MONTHLY=True 면 월세도) 매물 중 보증금 MAX_DEPOSIT 이하, 월세 MAX_RENT 이하만 골라
- 같은 매물(여러 중개사가 올린 동일 매물)은 하나로 합쳐 엑셀로 저장합니다.

실행: run.bat 더블클릭 (또는 python naver_rent.py)
"""
import math
import random
import re
import sys
import json
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import urlencode

from browser_util import connect, launch_browser

# ===== 조건 설정 (필요하면 숫자만 바꾸세요) =====
STATION_NAME = "운정중앙역"
STATION_LAT = 37.71661     # 위도
STATION_LON = 126.72835    # 경도
RADIUS_KM = 3.5            # 반경 (km)
MAX_DEPOSIT = 32000        # 보증금 상한 (만원) = 3억 2천
INCLUDE_MONTHLY = False    # 월세 매물도 넣으려면 True
MAX_RENT = 90              # 월세 상한 (만원), INCLUDE_MONTHLY=True 일 때만 사용
# 조회할 시·구는 좌표 주변에서 자동으로 찾습니다.
# 자동으로 못 찾는 경우를 대비해 꼭 넣을 지역코드를 여기에 적어 둘 수 있습니다 (비워도 됨).
EXTRA_REGION_CODES = [
    "4148000000",  # 파주시
    "4128700000",  # 고양시 일산서구 (반경이 넓으면 걸칠 수 있음)
]
ESTATE_TYPES = ["APT", "JGC"]   # 아파트, 재건축
NEARBY_EXTRA_KM = 1.5          # "단지목록" 시트에 반경 밖 단지도 이만큼 더 보여줌
# 접속 인증용 시작 화면 (아무 아파트 단지 화면이면 됨)
START_URL = "https://new.land.naver.com/complexes/113942?a=APT:ABYG:JGC&e=RETAIL"
# =============================================

BASE = "https://new.land.naver.com"
# 엑셀 링크 후보. 실행할 때 실제로 열어 보고 매물 가격이 보이는 첫 번째 형식을 씀
LINK_FORMATS = [
    ("네이버페이 부동산 매물", "https://fin.land.naver.com/articles/{no}"),
    ("네이버 부동산 모바일 매물", "https://m.land.naver.com/article/info/{no}"),
    ("네이버 부동산 단지", BASE + "/complexes/{cno}?a=APT:ABYG:JGC&e=RETAIL&articleNo={no}"),
]
# 열린 크롬 창 안에서 직접 요청 (브라우저 밖에서 따로 요청하면 네이버가 429로 거절함)
FETCH_JS = """async ([url, auth]) => {
  const r = await fetch(url, {headers: {"authorization": auth, "accept": "application/json"},
                              credentials: "include"});
  return [r.status, await r.text()];
}"""
HEADERS_COLUMNS = ["번호", "아파트 명", "거리(km)", "세대 수", "주차 대 수", "평 수", "층 수", "방향",
                   "보증금", "월세", "입주가능일", "중개사 이름", "휴대폰", "부동산 명", "링크"]


def log(msg):
    print(msg, flush=True)


def pause(a=1.5, b=3.0):
    time.sleep(random.uniform(a, b))


def distance_km(lat1, lon1, lat2, lon2):
    r = 6371.0
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp = p2 - p1
    dl = math.radians(lon2 - lon1)
    h = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(math.sqrt(h))


def parse_price(text):
    """'2억 5,000' -> 25000, '3억' -> 30000, '9,000' -> 9000 (단위: 만원)"""
    if text is None:
        return None
    if isinstance(text, (int, float)):
        return int(text)
    s = str(text).replace(",", "").replace(" ", "")
    if not s:
        return None
    total = 0
    m = re.search(r"(\d+)억", s)
    if m:
        total += int(m.group(1)) * 10000
        s = s[m.end():]
    m = re.search(r"(\d+)", s)
    if m:
        total += int(m.group(1))
    return total


def format_price(manwon):
    """25000 -> '2억 5,000', 30000 -> '3억', 9000 -> '9,000'"""
    if manwon is None:
        return ""
    eok, rest = divmod(int(manwon), 10000)
    if eok and rest:
        return f"{eok}억 {rest:,}"
    if eok:
        return f"{eok}억"
    return f"{rest:,}"


def to_pyeong(area1, area2):
    """공급면적(㎡)을 평으로 환산. 예: '33.0평 (공급 109㎡ / 전용 84㎡)'"""
    try:
        a1 = float(area1)
    except (TypeError, ValueError):
        return ""
    text = f"{a1 / 3.3058:.1f}평 (공급 {a1:g}㎡"
    try:
        text += f" / 전용 {float(area2):g}㎡)"
    except (TypeError, ValueError):
        text += ")"
    return text


def move_in_text(detail):
    """입주가능일: 'NOW' -> '즉시입주', '20270301' -> '2027.03.01'"""
    ymd = str(detail.get("moveInPossibleYmd") or "")
    if ymd.upper() == "NOW":
        text = "즉시입주"
    elif re.fullmatch(r"\d{8}", ymd):
        text = f"{ymd[:4]}.{ymd[4:6]}.{ymd[6:]}"
    elif re.fullmatch(r"\d{6}", ymd):
        text = f"{ymd[:4]}.{ymd[4:6]}"
    else:
        text = detail.get("moveInTypeName") or ""
    if detail.get("moveInDiscussionPossibleYN") == "Y" and "협의" not in text:
        text = (text + " (협의가능)").strip()
    return text


def first(d, *keys, default=""):
    for k in keys:
        if isinstance(d, dict) and d.get(k) not in (None, ""):
            return d[k]
    return default


class Naver:
    def __init__(self, page, debug_path):
        self.page = page
        self.token = None
        self.debug = open(debug_path, "w", encoding="utf-8")
        self.logged_paths = {}

    def write_debug(self, path, params, status, text):
        # 같은 종류의 요청은 처음 3번만 기록 (파일이 너무 커지지 않게)
        key = re.sub(r"\d{5,}", "N", path)
        self.logged_paths[key] = self.logged_paths.get(key, 0) + 1
        if self.logged_paths[key] > 3:
            return
        self.debug.write(json.dumps({"path": path, "params": params, "status": status,
                                     "response": text[:4000]}, ensure_ascii=False) + "\n")
        self.debug.flush()

    def capture_token(self):
        def on_request(req):
            auth = req.headers.get("authorization", "")
            if auth.startswith("Bearer ") and not self.token:
                self.token = auth

        self.page.on("request", on_request)
        log("네이버 부동산 접속 중...")
        if "new.land.naver.com/complexes" not in self.page.url:
            self.page.goto(START_URL, wait_until="domcontentloaded", timeout=60000)
        for _ in range(60):
            if self.token:
                break
            self.page.wait_for_timeout(500)
        if not self.token:
            # 페이지 소스에 토큰이 들어있는 경우 대비
            m = re.search(r'"token"\s*:\s*"([^"]+)"', self.page.content())
            if m:
                self.token = "Bearer " + m.group(1)
        if not self.token:
            raise RuntimeError("네이버 부동산 인증 토큰을 얻지 못했습니다. 잠시 후 다시 실행해 주세요.")
        log("접속 완료")

    def get(self, path, params=None, retries=4):
        url = BASE + path + ("?" + urlencode(params) if params else "")
        for attempt in range(retries):
            try:
                status, text = self.page.evaluate(FETCH_JS, [url, self.token])
            except Exception as e:
                status, text = 0, f"(요청 오류: {e})"
            self.write_debug(path, params, status, text)
            if status == 200:
                pause()
                try:
                    return json.loads(text)
                except ValueError:
                    log(f"  (응답 형식이 예상과 다름: {path})")
                    return None
            if status in (0, 401, 403, 429) or status >= 500:
                wait = 30 * (attempt + 1)
                log(f"  (네이버 응답 {status}, {wait}초 쉬었다가 다시 시도)")
                time.sleep(wait)
                if status == 401:
                    self.token = None
                    self.capture_token()
                continue
            log(f"  (요청 실패 {status}: {path})")
            return None
        return None

    # --- 단지 찾기 ---
    def collect_dongs(self, cortar_no, dongs, visited, depth=0):
        """지역 코드 아래로 내려가며 좌표 근처의 동(법정동) 목록을 모음"""
        if cortar_no in visited or depth > 3:
            return
        visited.add(cortar_no)
        data = self.get("/api/regions/list", {"cortarNo": cortar_no}) or {}
        for r in data.get("regionList", []):
            if r.get("cortarType") == "sec":
                dongs[r["cortarNo"]] = r
                continue
            lat, lon = r.get("centerLat"), r.get("centerLon")
            # 시·도는 넓게, 시·군·구는 반경 + 15km 안만 내려감
            limit = 80 if depth == 0 else RADIUS_KM + 15
            if lat and lon and distance_km(STATION_LAT, STATION_LON, float(lat), float(lon)) > limit:
                continue
            self.collect_dongs(r["cortarNo"], dongs, visited, depth + 1)

    def find_complexes(self):
        dongs, visited = {}, set()
        log("주변 시·구를 찾는 중...")
        self.collect_dongs("0000000000", dongs, visited)        # 전국 → 시·도 → 시·군·구 → 동
        for code in EXTRA_REGION_CODES:
            self.collect_dongs(code, dongs, visited, depth=1)
        log(f"조회할 동: {len(dongs)}곳")
        found = {}
        for dong in dongs.values():
            dlat, dlon = dong.get("centerLat"), dong.get("centerLon")
            # 동 중심이 역에서 너무 먼 곳(읍·면 등)은 건너뜀
            if dlat and dlon and distance_km(STATION_LAT, STATION_LON, float(dlat), float(dlon)) > RADIUS_KM + 4:
                continue
            for etype in ESTATE_TYPES:
                data = self.get("/api/regions/complexes",
                                {"cortarNo": dong["cortarNo"], "realEstateType": etype, "order": ""}) or {}
                for c in data.get("complexList", []):
                    lat, lon = c.get("latitude"), c.get("longitude")
                    if lat is None or lon is None:
                        continue
                    d = distance_km(STATION_LAT, STATION_LON, float(lat), float(lon))
                    if d <= RADIUS_KM + NEARBY_EXTRA_KM:
                        c["_distance"] = d
                        c["_inside"] = d <= RADIUS_KM
                        c["_dong"] = dong.get("cortarName", "")
                        found[str(c["complexNo"])] = c
        return sorted(found.values(), key=lambda c: c["_distance"])

    def complex_detail(self, complex_no):
        data = self.get(f"/api/complexes/{complex_no}", {"sameAddressGroup": "true"}) or {}
        return data.get("complexDetail", data)

    # --- 매물 ---
    def articles(self, complex_no):
        """(매물 목록, 조회 성공 여부)"""
        out, page_no, ok = [], 1, True
        while True:
            params = {
                "realEstateType": ":".join(ESTATE_TYPES),
                "tradeType": "B1:B2" if INCLUDE_MONTHLY else "B1",   # B1=전세, B2=월세
                "tag": "::::::::",
                "rentPriceMin": 0, "rentPriceMax": MAX_RENT,
                "priceMin": 0, "priceMax": MAX_DEPOSIT,
                "areaMin": 0, "areaMax": 900000000,
                "showArticle": "false",
                "sameAddressGroup": "true",     # 같은 매물 하나로 묶기
                "priceType": "RETAIL",
                "page": page_no,
                "complexNo": complex_no,
                "type": "list",
                "order": "rank",
            }
            data = self.get(f"/api/articles/complex/{complex_no}", params)
            if data is None:
                ok = False
                break
            out.extend(data.get("articleList", []))
            if not data.get("isMoreData") or page_no >= 30:
                break
            page_no += 1
        return out, ok

    def article_info(self, article_no, complex_no):
        """(매물 상세, 중개사 정보)"""
        data = self.get(f"/api/articles/{article_no}", {"complexNo": complex_no}) or {}
        return data.get("articleDetail", {}) or {}, data.get("articleRealtor", {}) or {}


def dedupe_key(a):
    # 같은 매물 합치기는 네이버의 sameAddressGroup 이 처리하므로,
    # 여기서는 같은 매물번호가 두 번 나오는 경우만 거름
    return a.get("articleNo")


def squash(text):
    return re.sub(r"[\s,]", "", str(text or ""))


def check_link(context, url, price_text):
    """새 탭으로 url 을 열어, 404 가 아니고 매물 가격이 화면에 보이는지 확인"""
    tab = context.new_page()
    try:
        tab.goto(url, wait_until="domcontentloaded", timeout=30000)
        try:
            tab.wait_for_load_state("networkidle", timeout=10000)
        except Exception:
            pass
        tab.wait_for_timeout(2000)
        body = squash(tab.inner_text("body"))
        final = tab.url
        ok = ("404" not in final and "찾을수없" not in body
              and squash(price_text) in body)
        return ok, final
    except Exception as e:
        return False, f"오류: {e}"
    finally:
        tab.close()


def choose_link_format(context, samples, debug):
    """매물 몇 개로 링크 형식을 시험해서, 모두 통과한 첫 번째 형식을 고름"""
    for name, tmpl in LINK_FORMATS:
        results = []
        for s in samples:
            url = tmpl.format(cno=s["cno"], no=s["no"])
            ok, final = check_link(context, url, s["price"])
            debug.write(json.dumps({"link_test": name, "url": url, "final": final, "ok": ok},
                                   ensure_ascii=False) + "\n")
            debug.flush()
            results.append(ok)
            if not ok:
                break
        log(f"  링크 확인 - {name}: {'성공' if all(results) else '실패'}")
        if results and all(results):
            return name, tmpl
    return None, None


def save_excel(rows, path, complex_rows=(), link_note=""):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "전월세"
    ws.append(HEADERS_COLUMNS)
    for i, r in enumerate(rows, 1):
        ws.append([i] + r)

    thin = Side(style="thin", color="BFBFBF")
    head_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in row:
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if cell.row > 1:
                cell.alignment = Alignment(vertical="center",
                                           horizontal="left" if cell.column in (2, 6, 12, 14) else "center")
    # 링크 열: 클릭하면 네이버 부동산 매물 화면이 열리도록
    link_col = len(HEADERS_COLUMNS)
    for row in ws.iter_rows(min_row=2, min_col=link_col, max_col=link_col):
        cell = row[0]
        if cell.value:
            cell.hyperlink = cell.value
            cell.value = "매물 보기"
            cell.font = Font(color="0563C1", underline="single")
    widths = [6, 26, 9, 9, 10, 30, 9, 9, 13, 9, 16, 12, 15, 24, 11]
    for idx, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    cs = wb.create_sheet("단지목록")
    cs.append(["아파트 명", "동", "역까지 거리(km)", "포함 여부", "세대 수", "조건 맞는 매물", "조회 결과"])
    for r in complex_rows:
        cs.append(r)
    for cell in cs[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
    for idx, w in enumerate([28, 10, 15, 12, 9, 15, 12], 1):
        cs.column_dimensions[get_column_letter(idx)].width = w
    cs.freeze_panes = "A2"
    cs.auto_filter.ref = cs.dimensions

    info = wb.create_sheet("조건")
    for line in [
        ("기준", f"{STATION_NAME} 반경 {RADIUS_KM:g}km"),
        ("보증금", f"{format_price(MAX_DEPOSIT)} 이하"),
        ("월세", f"{MAX_RENT}만원 이하 (전세 포함)" if INCLUDE_MONTHLY else "제외 (전세만)"),
        ("중복", "같은 매물은 하나로 합침"),
        ("평 수", "공급면적 ÷ 3.3058"),
        ("역 좌표", f"{STATION_LAT}, {STATION_LON}"),
        ("출처", "네이버 부동산"),
        ("링크", link_note),
        ("수집 시각", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("매물 수", len(rows)),
    ]:
        info.append(line)
    info.column_dimensions["A"].width = 12
    info.column_dimensions["B"].width = 60
    wb.save(path)


def main():
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        log("playwright 가 설치되지 않았습니다. run.bat 으로 실행해 주세요.")
        return 1

    out_dir = Path(__file__).resolve().parent
    out_path = out_dir / f"{STATION_NAME}_전월세_{datetime.now():%Y%m%d_%H%M}.xlsx"

    proc = launch_browser(out_dir, START_URL)
    with sync_playwright() as p:
        browser = connect(p)
        context = browser.contexts[0]
        page = context.pages[0] if context.pages else context.new_page()
        nv = Naver(page, out_dir / "debug_log.txt")
        nv.capture_token()

        all_complexes = nv.find_complexes()
        complexes = [c for c in all_complexes if c["_inside"]]
        log(f"반경 {RADIUS_KM:g}km 안 아파트 단지: {len(complexes)}곳")
        if not complexes:
            log("단지를 하나도 찾지 못했습니다. 같은 폴더의 debug_log.txt 를 보내주세요.")

        rows, seen, complex_rows = [], set(), []
        for c in all_complexes:
            if not c["_inside"]:
                complex_rows.append([c.get("complexName", ""), c["_dong"], round(c["_distance"], 2),
                                     "반경 밖", first(c, "totalHouseholdCount", "totalHouseHoldCount"), "", ""])
        for ci, c in enumerate(complexes, 1):
            cno = str(c["complexNo"])
            name = c.get("complexName", "")
            log(f"[{ci}/{len(complexes)}] {name} ({c['_distance']:.2f}km)")
            detail = nv.complex_detail(cno)
            households = first(detail, "totalHouseHoldCount", "totalHouseholdCount",
                               default=first(c, "totalHouseholdCount", "totalHouseHoldCount"))
            parking = first(detail, "parkingPossibleCount", "parkingCount")
            if parking == "":
                per = first(detail, "parkingCountByHousehold")
                try:
                    parking = round(float(per) * float(households))
                except (TypeError, ValueError):
                    parking = ""

            arts, ok = nv.articles(cno)
            kept = 0
            for a in arts:
                a["_complexNo"] = cno
                deposit = parse_price(a.get("dealOrWarrantPrc"))
                rent = parse_price(a.get("rentPrc")) or 0
                is_jeonse = (a.get("tradeTypeCode") == "B1") or (a.get("tradeTypeName") == "전세")
                if deposit is None or deposit > MAX_DEPOSIT or rent > MAX_RENT:
                    continue
                if not is_jeonse and not INCLUDE_MONTHLY:
                    continue
                key = dedupe_key(a)
                if key in seen:
                    continue
                seen.add(key)

                detail, realtor = nv.article_info(a.get("articleNo"), cno)
                phone = first(realtor, "cellPhoneNo", "representativeTelNo")
                rows.append([
                    name,
                    round(c["_distance"], 2),
                    households,
                    parking,
                    to_pyeong(a.get("area1"), a.get("area2")),
                    a.get("floorInfo", ""),
                    a.get("direction", ""),
                    format_price(deposit),
                    "전세" if is_jeonse else rent,
                    move_in_text(detail),
                    first(realtor, "representativeName"),
                    phone,
                    first(realtor, "realtorName", default=a.get("realtorName", "")),
                    {"cno": cno, "no": a.get("articleNo"), "price": format_price(deposit)},  # 링크는 마지막에 만듦
                ])
                kept += 1
            log(f"    조건에 맞는 매물 {kept}건" + ("" if ok else " (조회 실패)"))
            complex_rows.append([name, c["_dong"], round(c["_distance"], 2), "포함",
                                 households, kept, "성공" if ok else "실패"])

        # 링크 형식 확인 (매물 2개로 시험)
        link_note = ""
        if rows:
            log("엑셀 링크 확인 중...")
            name, tmpl = choose_link_format(context, [r[-1] for r in rows[:2]], nv.debug)
            if not tmpl:
                name, tmpl = LINK_FORMATS[0]
                link_note = f"{name} 형식 (자동 확인 실패 - 매물이 안 맞으면 알려주세요)"
            else:
                link_note = f"{name} 형식 (실행 시 확인됨)"
            for r in rows:
                r[-1] = tmpl.format(cno=r[-1]["cno"], no=r[-1]["no"])

        browser.close()
    try:
        proc.terminate()
    except Exception:
        pass

    complex_rows.sort(key=lambda r: r[2])
    save_excel(rows, out_path, complex_rows, link_note)
    log("")
    log(f"완료! 총 {len(rows)}건 저장: {out_path}")
    if not rows:
        log("매물이 0건입니다. 같은 폴더의 debug_log.txt 를 보내주세요.")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as e:  # 사용자에게 오류 내용을 그대로 보여주기
        import traceback
        traceback.print_exc()
        log("")
        log("오류가 발생했습니다. 위 내용을 복사해서 보내주세요.")
        sys.exit(1)
